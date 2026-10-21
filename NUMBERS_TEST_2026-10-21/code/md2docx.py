import re, sys
from docx import Document
from docx.shared import Pt

def add_runs(par, text):
    # **bold** and `code` → bold / monospace runs
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**"): r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"): r = par.add_run(part[1:-1]); r.font.name = "Courier New"
        else: par.add_run(part.replace("*", ""))

def convert(src, dst):
    doc = Document(); doc.styles["Normal"].font.size = Pt(11)
    lines = open(src, encoding="utf-8").read().splitlines(); i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells): rows.append(cells)
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    p = t.cell(r, c).paragraphs[0]; add_runs(p, cell)
                    if r == 0:
                        for run in p.runs: run.bold = True
            doc.add_paragraph(); continue
        m = re.match(r"(#+)\s+(.*)", ln)
        if m: doc.add_heading(m.group(2), level=min(len(m.group(1)), 3))
        elif ln.startswith("- "): add_runs(doc.add_paragraph(style="List Bullet"), ln[2:])
        elif re.match(r"\s{2,}- ", ln): add_runs(doc.add_paragraph(style="List Bullet 2"), ln.strip()[2:])
        elif ln.strip(): add_runs(doc.add_paragraph(), ln)
        i += 1
    doc.save(dst)

for src in sys.argv[1:]:
    convert(src, src[:-3] + ".docx"); print("wrote", src[:-3] + ".docx")
